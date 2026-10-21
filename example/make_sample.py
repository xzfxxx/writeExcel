#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成示例模板 example/template-sample.xlsx（带样式，可直接拿去改）。

运行：python example/make_sample.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "template-sample.xlsx"

THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def build() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "客户信息"

    # 标题行：A1:D1 合并
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = "客户信息登记表"
    title.font = Font(name="微软雅黑", size=16, bold=True, color="1F4E78")
    title.alignment = CENTER
    title.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.row_dimensions[1].height = 30

    # 表头行
    headers = ["姓名", "联系电话", "公司名称", "备注"]
    for idx, text in enumerate(headers, 1):
        cell = ws.cell(row=2, column=idx, value=text)
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = CENTER
        cell.border = BORDER

    # 数据区 3 行，先画好边框
    for row in range(3, 6):
        ws.row_dimensions[row].height = 22
        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    for col, width in enumerate((16, 22, 30, 34), 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A3"

    # 第二张表：演示多表写入
    ws2 = wb.create_sheet("备注说明")
    ws2["A1"] = "这张表也可以往里写，数据条目里加 s:'备注说明' 即可"
    ws2["A1"].font = Font(name="微软雅黑", size=10, italic=True, color="808080")
    ws2.column_dimensions["A"].width = 60

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"已生成示例模板：{path}")
